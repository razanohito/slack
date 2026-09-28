# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from data import COMPANIES, NARIKA, OTHERS, SWOT, ACTIONS, SURVEY_DATE

OUT = os.path.join(os.path.dirname(__file__), "..", "STEAM教育_競合比較_元データ.xlsx")
FONT = "Meiryo"
NAVY = "1F2A44"
YEL = "F5B700"
STEMON_FILL = PatternFill("solid", fgColor="FFF4CC")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="E8ECF3")
thin = Side(style="thin", color="C9CED8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(bold=False, color="222222", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(True, NAVY, 14)
    ws["A2"] = sub
    ws["A2"].font = f(False, "666666", 9)


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = f(True, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def body_row(ws, row, values, is_stemon=False):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = f(bold=(i <= 2 and is_stemon))
        cell.alignment = WRAP
        cell.border = BORDER
        if is_stemon:
            cell.fill = STEMON_FILL


wb = Workbook()

# ---------- 0. はじめに ----------
ws = wb.active
ws.title = "はじめに"
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 100
ws["A1"] = "STEAM教育・子ども向けプログラミング教室 競合比較（元データ）"
ws["A1"].font = f(True, NAVY, 16)
rows = [
    ("目的", "社内の戦略検討用。ステモンと主要競合の料金・レッスン内容・講師・教材・課題を網羅的に比較し、打ち手を検討する"),
    ("調査日", SURVEY_DATE),
    ("対象", "自社1＋競合13社（ロボット・ものづくり系7、プログラミング特化系3、オンライン・通信系3）＋参考：ナリカ（教材サプライヤー）"),
    ("調査方法", "Web検索結果（検索エンジンが表示した公式サイト・比較サイト・口コミサイト・求人サイトの要約）から収集"),
    ("⚠ 重要な注意", "調査環境の制約で各社公式ページの本文は直接閲覧できていない。数値は検索結果の要約ベースのため、社外利用・意思決定の前に「要確認リスト」シートの項目を公式ページで確認すること"),
    ("確度の表記", "公式＝運営会社・公式サイト由来／二次＝比較・口コミ・ブログ由来／推定＝本資料での試算／要確認＝出典が古い・食い違う・一次未確認"),
    ("料金の注意", "FC中心のブランドは教室ごとに料金が異なる。総額は「料金比較」シートの前提（C列）に基づく試算"),
    ("編集方法", "「料金比較」シートの青字セル（入会金・月謝・諸費用・教材費・授業時間）を書き換えると、総額・ステモン比・1時間単価・グラフ・サマリーが自動で再計算される"),
    ("シート構成", "サマリー／料金比較／レッスン・教材／講師／口コミ・課題／ポジショニング／ステモン戦略（SWOT・提言）／要確認リスト／参考_ナリカ・その他／出典"),
]
for i, (k, v) in enumerate(rows, 3):
    ws.cell(row=i, column=1, value=k).font = f(True, NAVY)
    c = ws.cell(row=i, column=2, value=v)
    c.font = f(bold=k.startswith("⚠"), color="B00020" if k.startswith("⚠") else "222222")
    c.alignment = WRAP
    ws.row_dimensions[i].height = 34

# ---------- 2. 料金比較 (build before summary so we know rows) ----------
wp = wb.create_sheet("料金比較")
title(wp, "料金比較（1年目・3年継続の総額試算）", "青字＝入力値（書き換え可）／黒字＝数式。金額は税込・円。教材・その他＝各年に一括でかかる費用（教材購入・更新料・大会参加費など）")
cols = ["分類", "社名", "計算の前提", "入会金", "月謝\n1年目", "月謝\n2年目", "月謝\n3年目", "月額\n諸費用", "教材・その他\n1年目", "教材・その他\n2年目", "教材・その他\n3年目",
        "1年目\n総額", "3年継続\n総額", "ステモン比\n(3年)", "月の授業\n時間(分)", "1時間あたり\n単価(1年目)", "月謝レンジ（参考）", "確度"]
widths = [13, 22, 46, 10, 10, 10, 10, 10, 12, 12, 12, 12, 12, 11, 10, 12, 26, 26]
HR = 4
header(wp, HR, cols, widths)
price_row = {}
first = HR + 1
stemon_row = first
for idx, c in enumerate(COMPANIES):
    r = first + idx
    price_row[c["key"]] = r
    st = c["key"] == "stemon"
    vals = [c["cat"], c["name"], c["price_assume"], c["entry"], c["m1"], c["m2"], c["m3"], c["misc"], c["mat1"], c["mat2"], c["mat3"],
            f"=D{r}+(E{r}+H{r})*12+I{r}",
            f"=D{r}+(E{r}+F{r}+G{r}+3*H{r})*12+I{r}+J{r}+K{r}",
            f"=M{r}/$M${stemon_row}",
            c["min_per_month"],
            f'=IF(O{r}="","—",(E{r}+H{r})/(O{r}/60))',
            c["fee_range"], c["price_conf"]]
    body_row(wp, r, vals, st)
    for col in range(4, 12):
        wp.cell(row=r, column=col).font = Font(name=FONT, color="0000FF", size=10, bold=st and False)
        wp.cell(row=r, column=col).number_format = '#,##0;(#,##0);"-"'
    wp.cell(row=r, column=15).font = Font(name=FONT, color="0000FF", size=10)
    for col in (12, 13, 16):
        wp.cell(row=r, column=col).number_format = '#,##0'
    wp.cell(row=r, column=13).font = f(True)
    wp.cell(row=r, column=14).number_format = '0%'
    wp.row_dimensions[r].height = 58
last = first + len(COMPANIES) - 1
n = last + 2
wp.cell(row=n, column=1, value="注記").font = f(True, NAVY)
notes = [
    "・Tech Kids School は新規受付停止中のため参考値。QUREOは入会金不明（0円で計算）。ロボ団はレンタル料・タブレット代を含まない。",
    "・クレファスは小4の追加キット、エジソンアカデミーは3年目の教材費が不明のため含まない（実際はさらに高い）。",
    "・1時間あたり単価＝（月謝＋月額諸費用）÷月の授業時間。授業時間が不明・通信型は「—」。",
    "・オンライン・通信系（Z会・デジタネ）は講師の対面指導がないため、単純な価格比較には注意。",
]
for i, t in enumerate(notes, 1):
    wp.cell(row=n + i, column=1, value=t).font = f(size=9, color="555555")

# chart: 3年総額
ch = BarChart()
ch.type = "bar"
ch.style = 10
ch.title = "3年継続の総額（試算、円）"
ch.y_axis.title = None
ch.x_axis.title = None
data = Reference(wp, min_col=13, min_row=HR, max_row=last)
cats = Reference(wp, min_col=2, min_row=first, max_row=last)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.legend = None
ch.height = 11
ch.width = 20
ch.y_axis.numFmt = '#,##0'
ch.y_axis.majorGridlines = None
ch.x_axis.scaling.orientation = "maxMin"
ch.series[0].graphicalProperties.solidFill = NAVY
ch.dataLabels = DataLabelList()
ch.dataLabels.showVal = True
ch.x_axis.delete = False
ch.y_axis.delete = False
wp.add_chart(ch, f"C{n + 7}")

# ---------- 1. サマリー ----------
wsum = wb.create_sheet("サマリー", 1)
title(wsum, "サマリー（全社の主要項目一覧）", "総額は「料金比較」シートと連動。脅威度＝ステモンにとっての競合度合い（高／中／低、定性評価）")
cols = ["分類", "社名", "運営会社", "対象", "規模", "運営形態", "月謝", "1年目総額\n(試算)", "3年総額\n(試算)", "教材（購入/貸与）", "講師（雇用・時給）", "強み（口コミで評価される点）", "主な弱み・課題", "脅威度"]
widths = [13, 22, 26, 18, 26, 18, 22, 11, 11, 30, 28, 40, 40, 8]
header(wsum, 4, cols, widths)
for idx, c in enumerate(COMPANIES):
    r = 5 + idx
    pr = price_row[c["key"]]
    st = c["key"] == "stemon"
    vals = [c["cat"], c["name"], c["operator"], c["target"], c["scale"], c["form"], c["fee_range"],
            f"='料金比較'!L{pr}", f"='料金比較'!M{pr}",
            f'{c["material"]}／{c["buy_lend"]}', f'{c["employment"]}／{c["wage"]}', c["good"], c["weakness"], c["threat"]]
    body_row(wsum, r, vals, st)
    for col in (8, 9):
        wsum.cell(row=r, column=col).number_format = '#,##0'
        wsum.cell(row=r, column=col).font = Font(name=FONT, color="008000", size=10, bold=st)
    t = wsum.cell(row=r, column=14)
    t.alignment = CENTER
    color = {"高": "C62828", "中": "E08A00", "低": "2E7D32"}.get(c["threat"], "222222")
    t.font = f(True, color)
    wsum.row_dimensions[r].height = 92

# ---------- 3. レッスン・教材 ----------
wl = wb.create_sheet("レッスン・教材")
title(wl, "レッスン内容・使用教材", "クラス人数は講師1人あたりの目安を含む")
cols = ["分類", "社名", "対象", "時間・回数", "クラス人数・講師比", "カリキュラム・コース構成", "大会・検定との接続", "振替", "保護者へのフィードバック", "使用教材", "購入／貸与・持ち帰り"]
widths = [13, 22, 18, 24, 22, 44, 34, 16, 26, 28, 26]
header(wl, 4, cols, widths)
for idx, c in enumerate(COMPANIES):
    r = 5 + idx
    body_row(wl, r, [c["cat"], c["name"], c["target"], c["lesson"], c["class_size"], c["curriculum"], c["contest"], c["furikae"], c["parent_fb"], c["material"], c["buy_lend"]], c["key"] == "stemon")
    wl.row_dimensions[r].height = 80

# ---------- 4. 講師 ----------
wt = wb.create_sheet("講師")
title(wt, "先生の質（募集要項・口コミ）", "時給は求人サイトの掲載例。FC中心のブランドは加盟事業者ごとに条件が異なる")
cols = ["分類", "社名", "雇用形態", "時給・給与", "応募条件", "研修制度", "講師1人あたり生徒数", "口コミ傾向（良い点／不満点）"]
widths = [13, 22, 24, 30, 24, 30, 18, 50]
header(wt, 4, cols, widths)
for idx, c in enumerate(COMPANIES):
    r = 5 + idx
    body_row(wt, r, [c["cat"], c["name"], c["employment"], c["wage"], c["qualification"], c["training"], c["ratio"], c["teacher_review"]], c["key"] == "stemon")
    wt.row_dimensions[r].height = 52

# ---------- 5. 口コミ・課題 ----------
wr = wb.create_sheet("口コミ・課題")
title(wr, "口コミ傾向・考えられる課題・ステモンの攻めどころ", "口コミはコエテコ・塾ナビ・オリコン・ブログ等の傾向を要約")
cols = ["分類", "社名", "褒められている点", "不満が多い点", "考えられる課題（弱み）", "ステモンが攻められるポイント", "脅威度"]
widths = [13, 22, 40, 40, 44, 48, 8]
header(wr, 4, cols, widths)
for idx, c in enumerate(COMPANIES):
    r = 5 + idx
    body_row(wr, r, [c["cat"], c["name"], c["good"], c["bad"], c["weakness"], c["attack"], c["threat"]], c["key"] == "stemon")
    wr.cell(row=r, column=7).alignment = CENTER
    wr.row_dimensions[r].height = 80

# ---------- 6. ポジショニング ----------
wpos = wb.create_sheet("ポジショニング")
title(wpos, "ポジショニング（価格 × 学びのスタイル）", "体験スコア：1＝画面内で完結するプログラミング／5＝実物を使ったものづくり・仕組み理解（定性評価、青字は書き換え可）")
cols = ["分類", "社名", "3年総額（試算）", "体験スコア(1-5)", "対象年齢の下限", "脅威度"]
widths = [13, 22, 16, 16, 16, 10]
header(wpos, 4, cols, widths)
for idx, c in enumerate(COMPANIES):
    r = 5 + idx
    pr = price_row[c["key"]]
    body_row(wpos, r, [c["cat"], c["name"], f"='料金比較'!M{pr}", c["exp_score"], c["min_age"], c["threat"]], c["key"] == "stemon")
    wpos.cell(row=r, column=3).number_format = '#,##0'
    wpos.cell(row=r, column=3).font = Font(name=FONT, color="008000", size=10)
    wpos.cell(row=r, column=4).font = Font(name=FONT, color="0000FF", size=10)
wpos.cell(row=5 + len(COMPANIES) + 1, column=1, value="読み取り：「3年40〜50万円 × ものづくり寄り」にステモン・ヒューマンアカデミー・エジソンアカデミー・プログラボ・STEAM Campusが集中。価格帯内での差別化軸（講師品質・成果の見える化・年中からの一貫性）が重要。").font = f(size=9, color="555555")

# ---------- 7. ステモン戦略 ----------
wsw = wb.create_sheet("ステモン戦略")
title(wsw, "ステモンのSWOTと提言", "競合比較の結果から導出（社内検討用のたたき台）")
wsw.column_dimensions["A"].width = 16
wsw.column_dimensions["B"].width = 60
wsw.column_dimensions["C"].width = 16
wsw.column_dimensions["D"].width = 60
labels = [("S", "強み（Strength）"), ("W", "弱み（Weakness）"), ("O", "機会（Opportunity）"), ("T", "脅威（Threat）")]
r = 4
for (k1, l1), (k2, l2) in [(labels[0], labels[1]), (labels[2], labels[3])]:
    for col, lab in ((1, l1), (3, l2)):
        cell = wsw.cell(row=r, column=col, value=lab)
        cell.font = f(True, "FFFFFF")
        cell.fill = HEAD_FILL
        wsw.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
    for i in range(5):
        for col, k in ((1, k1), (3, k2)):
            cell = wsw.cell(row=r + 1 + i, column=col, value=f"{k}{i + 1}")
            cell.font = f(True, NAVY)
            cell.alignment = CENTER
            cell.border = BORDER
            c2 = wsw.cell(row=r + 1 + i, column=col + 1, value=SWOT[k][i])
            c2.font = f()
            c2.alignment = WRAP
            c2.border = BORDER
        wsw.row_dimensions[r + 1 + i].height = 32
    r += 7
r += 1
for i, h in enumerate(["領域", "打ち手", "優先度", "期間の目安"], 1):
    cell = wsw.cell(row=r, column=i, value=h)
    cell.font = f(True, "FFFFFF")
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
for a in ACTIONS:
    r += 1
    for i, v in enumerate(a, 1):
        cell = wsw.cell(row=r, column=i, value=v)
        cell.font = f(bold=(i == 1))
        cell.alignment = WRAP if i == 2 else CENTER
        cell.border = BORDER
    wsw.row_dimensions[r].height = 36

# ---------- 8. 要確認リスト ----------
wq = wb.create_sheet("要確認リスト")
title(wq, "要確認リスト（公式ページ等で確認が必要な項目）", "確認したら「状態」を更新し、該当シートの値を修正する")
cols = ["No", "対象", "確認項目", "現在の値・状況", "確認先", "優先度", "状態"]
widths = [5, 20, 36, 44, 34, 8, 10]
header(wq, 4, cols, widths)
wq.freeze_panes = "A5"
checks = [
    ("ステモン", "入会金・施設管理費の正しい金額", "入会金16,500円 or 5,500円／施設費880円 or 2,200円で情報が食い違う", "社内（料金表）", "高"),
    ("ステモン", "プログラミング＆ロボティクス（小4〜）の月謝", "未取得", "社内（料金表）", "高"),
    ("ステモン", "振替ルール・保護者レポートの有無・教材製品名", "教室ごとに異なる（要確認）", "社内", "中"),
    ("STEAM Campus", "コース別の月謝、設備費・指導関連費、月の回数、教室総数", "月謝は二次情報の平均値11,100円", "steamcampus.jp／各運営会社サイト", "高"),
    ("ヒューマンアカデミー", "現在のキット代・進級時の追加キット額", "33,000円（旧31,350円）、追加は約1〜2.6万円と食い違い", "kids.athuman.com/robo/", "中"),
    ("ロボ団", "コース別月謝、教材レンタル料・タブレット代の有無、最新教室数、講師時給", "月謝9,800〜15,800円で出典が食い違う", "robo-done.com", "高"),
    ("クレファス", "コース別月謝、追加キット額、定員", "3年総額が別サイト試算（年約15万円）と大きく食い違う", "crefus.jp／近隣教室", "高"),
    ("エジソンアカデミー", "最新教室数、3年目教材費、講師募集条件", "教室数は2021年頃の約850〜900", "artec-kk.co.jp", "中"),
    ("プログラボ", "2026年時点の校数、保護者向けフィードバック制度", "84校（2024年9月）", "proglab.education", "中"),
    ("タミヤロボットスクール", "最新教室数、対象学年の下限、クラス人数、3年目教材費", "約100教室（古い）", "tamiya-robotschool.com", "低"),
    ("LITALICOワンダー", "入会金16,500円、生徒数約5,050名の時点", "入会金は二次情報", "wonder.litalico.jp", "低"),
    ("Tech Kids School", "新規受付停止の開始時期・理由・再開予定", "公式に「受付停止」の記載", "techkidsschool.jp", "中"),
    ("QUREO", "入会金・教室ごとの月謝、最新教室数と基準日、加盟条件", "入会金不明（0円で試算）", "qureo.jp／近隣加盟塾", "高"),
    ("Z会", "みらい講座の新料金（6,200円の内訳）、ステージ2・3の額", "新旧で表記が異なる", "zkai.co.jp", "低"),
    ("デジタネ", "対象年齢、提携教室数、講師体制", "未取得", "digitane.jp", "低"),
    ("CodeCampKIDS", "教材費、月の回数、定員、振替", "未取得", "codecampkids.jp", "低"),
    ("全社", "講師1人あたり生徒数の公式値、口コミ評価点（コエテコ・塾ナビ）の最新値", "一部のみ取得", "各社サイト・口コミサイト", "中"),
]
for i, row in enumerate(checks, 1):
    r = 4 + i
    body_row(wq, r, [i, *row, "未確認"])
    wq.cell(row=r, column=1).alignment = CENTER
    wq.cell(row=r, column=6).alignment = CENTER
    wq.cell(row=r, column=7).alignment = CENTER
    wq.row_dimensions[r].height = 34

# ---------- 9. 参考_ナリカ・その他 ----------
wn = wb.create_sheet("参考_ナリカ・その他")
title(wn, NARIKA["name"], NARIKA["summary"])
wn.column_dimensions["A"].width = 22
wn.column_dimensions["B"].width = 100
wn.column_dimensions["C"].width = 36
header(wn, 4, ["項目", "内容"])
wn.freeze_panes = None
for i, (k, v) in enumerate(NARIKA["rows"], 5):
    body_row(wn, i, [k, v])
    wn.row_dimensions[i].height = 32
r = 5 + len(NARIKA["rows"]) + 1
wn.cell(row=r, column=1, value="その他 注視すべき教室").font = f(True, NAVY, 12)
header_r = r + 1
for i, h in enumerate(["名称", "概要", "URL"], 1):
    cell = wn.cell(row=header_r, column=i, value=h)
    cell.font = f(True, "FFFFFF")
    cell.fill = HEAD_FILL
for i, o in enumerate(OTHERS, header_r + 1):
    body_row(wn, i, list(o))

# ---------- 10. 出典 ----------
wsrc = wb.create_sheet("出典")
title(wsrc, "出典一覧", f"確認日：{SURVEY_DATE}（すべて検索結果の要約で確認。本文の直接閲覧は未実施）")
header(wsrc, 4, ["社名", "URL"], [30, 110])
wsrc.freeze_panes = "A5"
r = 5
for c in COMPANIES + [dict(name=NARIKA["name"], sources=NARIKA["sources"], key="narika")]:
    for u in c["sources"]:
        body_row(wsrc, r, [c["name"], u], c.get("key") == "stemon")
        wsrc.cell(row=r, column=2).hyperlink = u
        wsrc.cell(row=r, column=2).font = Font(name=FONT, color="1155CC", underline="single", size=10)
        r += 1

for sh in wb.worksheets:
    sh.sheet_view.zoomScale = 90
wb.save(OUT)
print("saved", os.path.abspath(OUT))
